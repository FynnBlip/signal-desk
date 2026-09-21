# -*- coding: utf-8 -*-
"""07 闭环大脑 —— 降噪 DUT 迭代闭环（闭环一，人在环）。

流程：配置矩阵 → 跑一轮（归因 + 条件晋级判定）→ LLM/规则给建议 → 人在环确认 → 执行。
数据红线：分数全部来自 audiobox 真实推理；无解就写无解，不造晋级假象。
"""

import csv
import datetime
import hashlib
import json
import math
import os
import re
import threading
import uuid
from contextlib import contextmanager
from functools import wraps
from pathlib import Path

import numpy as np
from scipy import stats

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
DATA_DIR = PROJECT_ROOT / "data"
LOOP_DIR = DATA_DIR / "loop"
LOOP_STATUS = LOOP_DIR / "loop_status.json"
HISTORY_ARCHIVE_DIR = LOOP_DIR / "archive"
HISTORY_ARCHIVE_FILE = HISTORY_ARCHIVE_DIR / "loop_history.json"
HISTORY_ARCHIVE_AFTER_DAYS = 7
MATRIX_DIR = DATA_DIR / "matrix"
EXAM_DIR = DATA_DIR / "exam"
DENOISED_DIR = DATA_DIR / "denoised"
LISTENING_LOG = LOOP_DIR / "listening_v1.jsonl"
ACTIVE_COHORT = DATA_DIR / "voice" / "active_cohort.json"
VOICE_COHORTS_DIR = DATA_DIR / "voice" / "cohorts"
BENCHMARK_ROOT = LOOP_DIR / "benchmarks"
_LISTENING_LOCK = threading.Lock()
_STATUS_LOCK = threading.RLock()
_CACHE_LOCK = threading.Lock()


def _load_env_file():
    env_path = PROJECT_ROOT / ".env"
    if not env_path.exists():
        return
    for line in env_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


_load_env_file()


def _load_module(path):
    import importlib.util

    spec = importlib.util.spec_from_file_location(path.stem, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


denoise = _load_module(PROJECT_ROOT / "modules" / "05_denoise" / "denoise.py")
evaluate = _load_module(PROJECT_ROOT / "modules" / "06_evaluate" / "evaluate.py")
channel = _load_module(PROJECT_ROOT / "modules" / "04_channel" / "channel.py")

DEEPSEEK_MODEL = os.environ.get("DEEPSEEK_MODEL", "deepseek-v4-flash-vision-exp")
DEEPSEEK_BASE_URL = os.environ.get("DEEPSEEK_BASE_URL", "https://api.deepseek.com/anthropic")
DEEPSEEK_API_KEY = os.environ.get("DEEPSEEK_API_KEY", "")

VERSION_CANDIDATES = [
    {"id": "v1", "model": "noisereduce", "strength": 1.0},
    {"id": "v2", "model": "noisereduce", "strength": 0.8},
    {"id": "v3", "model": "noisereduce", "strength": 0.6},
    {"id": "v4", "model": "noisereduce", "strength": 0.4},
]

# 产品评审已锁定：场景按业务重要性固定加权；声线簇不另行等权，
# 每个真实音色等权后，10/7/5/4 的簇规模会自然形成簇权重。
SCENE_WEIGHTS = {
    "NPARK": 0.25,
    "OOFFICE": 0.25,
    "PCAFETER": 0.20,
    "PRESTO": 0.15,
    "STRAFFIC": 0.10,
    "TMETRO": 0.05,
}

DEFAULT_CONFIG = {
    "matrix": {"source": "golden", "exam": None, "noise_scenes": []},
}

GOLDEN_BENCHMARK = {
    "id": "golden-v1",
    "label": "Golden Benchmark v1",
    "ready": False,
    "probe": "固定探针语料",
    "voice_set": "26 个参考音色 · C1–C4",
    "scene_set": "固定噪声场景与 SNR",
    "evaluator": "AudioBox PQ · Gate v2",
}

# 这是工程 guardrail，不伪装成统计学显著性标准。p 值只作为旁证展示；
# 样本数、均值、回退比例与最差常规场景共同决定是否允许晋级。
GATE_POLICY = {
    "version": "denoise-pq-gate-v2",
    "min_regular_samples": 5,
    "min_regular_mean_delta": 0.05,
    "max_regular_retreat_ratio": 0.40,
    "min_regular_worst_delta": -0.25,
}

# Cache identity schema. Old golden-v1-* directories stay on disk as history.
ASSET_IDENTITY_SCHEMA = 2
DENOISE_IMPL_VERSION = "05-process-v1"
EVALUATE_IMPL_VERSION = "06-audiobox-v1"
_CONTENT_HASH_MEMO = {}


def list_exams():
    if not EXAM_DIR.exists():
        return []
    return [{"stem": p.stem, "name": p.name} for p in sorted(EXAM_DIR.glob("*.wav"))]


def list_scenes():
    return [{"id": s["id"], "label": s["label"], "snr_db": s["snr_db"]} for s in channel.NOISE_SCENES]


def list_versions():
    """返回版本候选（demo：NoiseReduce 按真实降噪比例 prop_decrease 分档）。"""
    return [
        {
            "id": v["id"],
            "model": v["model"],
            "strength": v["strength"],
            "label": f"{v['id']} · {v['model']} · 降噪比例 {v['strength']:g}",
        }
        for v in VERSION_CANDIDATES
    ]


def get_version(vid):
    for v in VERSION_CANDIDATES:
        if v["id"] == vid:
            return v
    return None


def _load_active_cohort():
    """Resolve the immutable 26-voice input set and its anchored C1-C4 labels."""
    if not ACTIVE_COHORT.exists():
        raise RuntimeError("尚未激活固定音色 Cohort，请先在声线基线页激活一个版本。")
    pointer = json.loads(ACTIVE_COHORT.read_text(encoding="utf-8"))
    cohort_id = Path(str(pointer.get("cohort_id") or "")).name
    cohort_dir = VOICE_COHORTS_DIR / cohort_id
    manifest_path = cohort_dir / "manifest.json"
    result_path = cohort_dir / "cluster_result.json"
    if not manifest_path.exists() or not result_path.exists():
        raise RuntimeError(f"固定音色 Cohort 不完整：{cohort_id}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    result = json.loads(result_path.read_text(encoding="utf-8"))
    manifest_items = {
        str(item.get("filename")): item for item in manifest.get("items") or []
        if item.get("status") == "ok" and item.get("filename")
    }
    voices = []
    for item in result.get("files") or []:
        filename = str(item.get("file") or "")
        audio = cohort_dir / "audio" / filename
        source = manifest_items.get(filename) or {}
        if not audio.exists() or item.get("cluster") is None:
            continue
        voices.append({
            "file": filename,
            "path": str(audio),
            "voice_id": item.get("voice_id") or source.get("voice_id"),
            "voice_name": item.get("voice_name") or source.get("voice_name"),
            "cluster_id": int(item["cluster"]) + 1,
            "sha256": source.get("sha256"),
        })
    voice_ids = [str(item.get("voice_id") or "").strip() for item in voices]
    filenames = [item["file"] for item in voices]
    clusters = {item["cluster_id"] for item in voices}
    if len(voices) != 26 or len(set(voice_ids)) != 26 or "" in voice_ids or len(set(filenames)) != 26:
        raise RuntimeError(f"Golden Benchmark 要求 26 个有效音色，当前 Cohort 只有 {len(voices)} 个。")
    if clusters != {1, 2, 3, 4}:
        raise RuntimeError(f"Golden Benchmark 要求固定 C1–C4，当前簇为 {sorted(clusters)}。")
    return {"cohort_id": cohort_id, "cohort_dir": str(cohort_dir), "voices": voices}


def _sha256_file(path, *, memoize=True):
    """Content hash. Memoize by path+size+mtime so homepage polls do not reread noise."""
    path = Path(path)
    stat = path.stat()
    key = (str(path.resolve()), stat.st_size, stat.st_mtime_ns)
    if memoize and key in _CONTENT_HASH_MEMO:
        return _CONTENT_HASH_MEMO[key]
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    hexdigest = digest.hexdigest()
    if memoize:
        _CONTENT_HASH_MEMO[key] = hexdigest
    return hexdigest


def _noise_identity_for_scenes(scenes):
    rows = []
    for scene in scenes:
        files = channel.scene_noise_files(scene)
        rows.append({
            "id": scene["id"],
            "snr_db": scene["snr_db"],
            "dir": scene["dir"],
            "files": [{"name": path.name, "sha256": _sha256_file(path)} for path in files],
        })
    return rows


def _golden_benchmark_context(config=None):
    cohort = _load_active_cohort()
    matrix = (config or {}).get("matrix") or {}
    scene_ids = matrix.get("noise_scenes") or [item["id"] for item in channel.NOISE_SCENES]
    scenes = [item for item in channel.NOISE_SCENES if item["id"] in scene_ids]
    payload = {
        "schema_version": ASSET_IDENTITY_SCHEMA,
        "benchmark_id": GOLDEN_BENCHMARK["id"],
        "cohort_id": cohort["cohort_id"],
        "voices": [{"file": item["file"], "sha256": item.get("sha256"), "cluster_id": item["cluster_id"]} for item in cohort["voices"]],
        "scenes": _noise_identity_for_scenes(scenes),
        "channel": {
            "impl_version": channel.CHANNEL_IMPL_VERSION,
            "bandwidth": "wideband",
            "bitrate_kbps": 16,
            "cbr": False,
            "seed": 42,
        },
        "versions": VERSION_CANDIDATES,
    }
    fingerprint = hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()[:16]
    root = BENCHMARK_ROOT / f"golden-v1-{fingerprint}"
    return {**cohort, "scenes": scenes, "fingerprint": fingerprint, "root": root, "payload": payload}


def _golden_matrix_filename(voice, scene):
    stem = Path(voice["file"]).stem
    return f"{stem}__{scene['id']}_snr{scene['snr_db']}dB_opus_wideband_16k_vbr_seed42.wav"


def prepare_golden_matrix(config=None, progress_cb=None):
    """Materialise and reuse the fixed 26 voices × selected scenes matrix."""
    blockers = channel.benchmark_noise_blockers()
    if blockers:
        raise RuntimeError("；".join(blockers))
    context = _golden_benchmark_context(config)
    matrix_dir = context["root"] / "matrix"
    matrix_dir.mkdir(parents=True, exist_ok=True)
    total = len(context["voices"]) * len(context["scenes"])
    done = 0
    samples = []
    for voice in context["voices"]:
        for scene in context["scenes"]:
            target = matrix_dir / _golden_matrix_filename(voice, scene)
            if not target.exists():
                generated = channel.run(
                    voice["path"], noise_scenes=[scene["id"]], bandwidth="wideband",
                    bitrate_kbps=16, cbr=False, seed=42, output_dir=matrix_dir,
                )
                run = next((item for item in generated.get("runs") or [] if item.get("status") == "ok"), None)
                if not run or not Path(run.get("degraded") or "").exists():
                    message = next((item.get("message") for item in generated.get("runs") or [] if item.get("status") == "error"), "信道生成失败")
                    raise RuntimeError(f"{voice['voice_name']} · {scene['label']}：{message}")
                target = Path(run["degraded"])
            done += 1
            if progress_cb:
                progress_cb(done, total, f"准备固定测试矩阵 · {voice['voice_name']} · {scene['label']}")
            samples.append({
                "scene_id": scene["id"], "noise_label": scene["label"], "snr_db": scene["snr_db"],
                "clean": voice["path"], "degraded": str(target), "stem": target.stem,
                "voice_id": voice["voice_id"], "voice_name": voice["voice_name"],
                "cluster_id": voice["cluster_id"], "cohort_id": context["cohort_id"],
                "denoised_root": str(context["root"] / "denoised"),
            })
    metadata = {
        **context["payload"], "fingerprint": context["fingerprint"],
        "matrix_dir": str(matrix_dir), "sample_count": len(samples),
        "created_or_verified_at": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    metadata_path = context["root"] / "metadata.json"
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return samples, context


def _noise_scene_from_stem(stem):
    m = re.search(r"__([A-Z]+)_snr", stem)
    return m.group(1) if m else None


def collect_matrix(config, progress_cb=None):
    """按配置确定回归矩阵：固定 exam + 放开/固定的噪声场景。"""
    if ((config.get("matrix") or {}).get("source") or "golden") == "golden":
        samples, _ = prepare_golden_matrix(config, progress_cb=progress_cb)
        return samples
    exam = (config.get("matrix") or {}).get("exam")
    scenes = (config.get("matrix") or {}).get("noise_scenes") or [s["id"] for s in channel.NOISE_SCENES]
    meta = evaluate.load_channel_meta()
    samples = []
    for degraded in sorted(MATRIX_DIR.glob("*.wav")):
        stem = degraded.stem
        scene = _noise_scene_from_stem(stem)
        if scene not in scenes:
            continue
        if exam and not stem.startswith(exam + "__"):
            continue
        clean = EXAM_DIR / f"{stem.split('__')[0]}.wav"
        if not clean.exists():
            continue
        m = meta.get(degraded.name, {})
        samples.append(
            {
                "scene_id": scene,
                "noise_label": m.get("noise_label", scene),
                "snr_db": m.get("snr_db", ""),
                "clean": str(clean),
                "degraded": str(degraded),
                "stem": stem,
            }
        )
    return samples


def _dut_identity(version):
    model = version["model"]
    if model == "noisereduce":
        provenance = denoise._provider_provenance(model)
    elif model in denoise.MODEL_FILES:
        path = denoise.MODEL_DIR / denoise.MODEL_FILES[model]
        provenance = {
            "provider_version": "onnx",
            "model_artifact": path.name,
            "model_artifact_sha256": _sha256_file(path) if path.is_file() else "",
        }
    else:
        provenance = {"provider_version": "unknown", "model_artifact": model, "model_artifact_sha256": ""}
    payload = {
        "schema_version": ASSET_IDENTITY_SCHEMA,
        "impl_version": DENOISE_IMPL_VERSION,
        "model": model,
        "strength": version.get("strength"),
        "provider_version": provenance.get("provider_version") or "",
        "model_artifact": provenance.get("model_artifact") or "",
        "model_artifact_sha256": provenance.get("model_artifact_sha256") or "",
    }
    fingerprint = hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()[:12]
    return {**payload, "fingerprint": fingerprint}


def _denoised_path(sample, version):
    model = version["model"]
    suffix = "" if version["strength"] >= 1.0 else f"__s{int(round(version['strength'] * 100))}"
    root = Path(sample.get("denoised_root") or DENOISED_DIR)
    identity = _dut_identity(version)
    return root / model / identity["fingerprint"] / f"{sample['stem']}__{model}{suffix}.wav"


def _denoise_identity_sidecar(output):
    output = Path(output)
    return output.with_name(output.name + ".identity.json")


def _expected_denoise_identity(sample, version):
    degraded = Path(sample["degraded"])
    return {
        "schema_version": ASSET_IDENTITY_SCHEMA,
        "input_sha256": _sha256_file(degraded) if degraded.is_file() else "",
        "dut": _dut_identity(version),
    }


def _denoise_output_reusable(sample, version):
    output = _denoised_path(sample, version)
    sidecar = _denoise_identity_sidecar(output)
    if not output.is_file() or not sidecar.is_file():
        return False
    try:
        recorded = json.loads(sidecar.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return False
    return recorded == _expected_denoise_identity(sample, version)


def _write_json_atomic(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(tmp, path)
    finally:
        if tmp.exists():
            tmp.unlink()


def ensure_denoised(samples, version, progress_cb=None):
    """补齐某个版本在矩阵上的降噪产物（身份匹配才复用）。"""
    missing = []
    for s in samples:
        if not _denoise_output_reusable(s, version):
            missing.append(s)
    if not missing:
        if progress_cb:
            progress_cb(len(samples), len(samples), f"降噪产物复用（{version['id']}）· {len(samples)} 条")
        return {"created": 0}
    if progress_cb:
        progress_cb(0, len(missing), f"降噪补齐（{version['id']}）…")
    for i, s in enumerate(missing, 1):
        output = _denoised_path(s, version)
        output.parent.mkdir(parents=True, exist_ok=True)
        expected = _expected_denoise_identity(s, version)
        temp_output = output.with_name(f".{output.stem}.{uuid.uuid4().hex}.tmp.wav")
        try:
            result = denoise.process(
                s["degraded"], temp_output,
                {"model": version["model"], "strength": version["strength"]},
            )
            if result.get("status") != "ok":
                raise RuntimeError(result.get("message") or f"{version['id']} 降噪失败")
            os.replace(temp_output, output)
            _write_json_atomic(_denoise_identity_sidecar(output), expected)
        finally:
            if temp_output.exists():
                temp_output.unlink()
        if progress_cb:
            progress_cb(i, len(missing), f"降噪补齐（{version['id']}）…")
    return {"created": len(missing)}


def judge(baseline_pqs, candidate_pqs, flags):
    """工程晋级判定：均值上行不能掩盖多数场景回退。"""
    reg_base = [b for b, f in zip(baseline_pqs, flags) if f]
    reg_cand = [c for c, f in zip(candidate_pqs, flags) if f]
    ext_base = [b for b, f in zip(baseline_pqs, flags) if not f]
    ext_cand = [c for c, f in zip(candidate_pqs, flags) if not f]

    reg_delta = float(np.mean(np.array(reg_cand) - np.array(reg_base))) if reg_base else 0.0
    ext_delta = float(np.mean(np.array(ext_cand) - np.array(ext_base))) if ext_base else 0.0

    reg_p = None
    if len(reg_base) >= 2:
        try:
            _, reg_p = stats.ttest_rel(reg_cand, reg_base)
            reg_p = float(reg_p)
            if math.isnan(reg_p):
                reg_p = None
        except Exception:
            reg_p = None

    reg_deltas = [float(c - b) for b, c, f in zip(baseline_pqs, candidate_pqs, flags) if f]
    extreme_bypass = bool(ext_base) and ext_delta < -0.05
    reg_retreat = int(sum(1 for b, c, f in zip(baseline_pqs, candidate_pqs, flags) if f and c < b))
    retreat_ratio = reg_retreat / len(reg_deltas) if reg_deltas else 1.0
    worst_delta = min(reg_deltas) if reg_deltas else None
    checks = [
        {
            "id": "sample_count",
            "label": "常规样本数",
            "pass": len(reg_deltas) >= GATE_POLICY["min_regular_samples"],
            "actual": len(reg_deltas),
            "threshold": f">={GATE_POLICY['min_regular_samples']}",
        },
        {
            "id": "mean_delta",
            "label": "常规平均 ΔPQ",
            "pass": reg_delta >= GATE_POLICY["min_regular_mean_delta"],
            "actual": round(reg_delta, 4),
            "threshold": f">={GATE_POLICY['min_regular_mean_delta']:+.2f}",
        },
        {
            "id": "retreat_ratio",
            "label": "常规场景回退比例",
            "pass": retreat_ratio <= GATE_POLICY["max_regular_retreat_ratio"],
            "actual": round(retreat_ratio, 4),
            "threshold": f"<={GATE_POLICY['max_regular_retreat_ratio']:.0%}",
        },
        {
            "id": "worst_delta",
            "label": "最差常规场景 ΔPQ",
            "pass": worst_delta is not None and worst_delta >= GATE_POLICY["min_regular_worst_delta"],
            "actual": round(worst_delta, 4) if worst_delta is not None else None,
            "threshold": f">={GATE_POLICY['min_regular_worst_delta']:+.2f}",
        },
    ]
    regular_pass = all(item["pass"] for item in checks)

    if regular_pass:
        if extreme_bypass:
            verdict = "条件晋级"
        elif reg_retreat == 0:
            verdict = "全绿晋级"
        else:
            verdict = f"整体晋级（{reg_retreat} 场景回退）"
    else:
        verdict = "回滚"

    return {
        "policy_version": GATE_POLICY["version"],
        "policy": dict(GATE_POLICY),
        "gate_checks": checks,
        "gate_reasons": [item["label"] for item in checks if not item["pass"]],
        "n_regular": len(reg_base),
        "n_extreme": len(ext_base),
        "regular_retreat": reg_retreat,
        "regular_retreat_ratio": round(retreat_ratio, 4),
        "regular_worst_delta": round(worst_delta, 4) if worst_delta is not None else None,
        "gate_status": "conditional" if regular_pass and extreme_bypass else "promoted" if regular_pass else "rejected",
        "promoted": regular_pass and not extreme_bypass,
        "conditional": regular_pass and extreme_bypass,
        "baseline_regular_pq": round(float(np.mean(reg_base)), 4) if reg_base else None,
        "candidate_regular_pq": round(float(np.mean(reg_cand)), 4) if reg_cand else None,
        "regular_delta": round(reg_delta, 4),
        "regular_p": round(reg_p, 6) if reg_p is not None else None,
        "extreme_delta": round(ext_delta, 4),
        "extreme_bypass": extreme_bypass,
        "verdict": verdict,
    }


def rule_decide(judge_report, config, evaluated):
    """规则兜底决策：新版本整体上行 → accept；回滚 → 迭代下一版 / 无解。"""
    candidate = config["candidate"]
    pool = [v["id"] for v in list_versions()]
    if judge_report.get("conditional"):
        return {
            "action": "accept_conditional",
            "version": candidate,
            "hypothesis": f"新版本 {candidate} 常规曲线上行，但极端噪声回退；仅可在明确知悉风险后条件晋级",
            "reason": judge_report["verdict"],
            "conditions": ["极端噪声场景存在回退", "该版本不得记录为无条件晋级"],
        }
    if judge_report.get("promoted"):
        return {
            "action": "accept",
            "version": candidate,
            "hypothesis": f"新版本 {candidate} 曲线整体上行（常规 Δ={judge_report['regular_delta']:+.3f}），验收晋级",
            "reason": judge_report["verdict"],
        }
    untried = [v for v in pool if v not in set(evaluated)]
    if untried:
        version = untried[0]
        return {
            "action": "iterate",
            "version": version,
            "hypothesis": f"新版本 {candidate} 曲线未整体上行，回滚；下一版 {version} 可针对薄弱场景定向调参",
            "reason": f"{candidate} 常规工况 Δ={judge_report['regular_delta']:+.3f}，未晋级，继续迭代",
        }
    return {
        "action": "no_solution",
        "version": None,
        "hypothesis": "已无可继续迭代的版本，诚实收口",
        "reason": "各版本均无整体正收益，保持已验收版本；本地不训练，如需继续需人工引入新模型/新训练版本",
    }


def allowed_actions(judge_report, evaluated):
    """Return the only actions an advisor may recommend for this gate result."""
    pool = [v["id"] for v in list_versions()]
    untried = [v for v in pool if v not in set(evaluated)]
    if judge_report.get("conditional"):
        return {"actions": ["accept_conditional"], "iterate_versions": []}
    if judge_report.get("promoted"):
        return {"actions": ["accept"], "iterate_versions": []}
    if untried:
        return {"actions": ["iterate"], "iterate_versions": untried}
    return {"actions": ["no_solution"], "iterate_versions": []}


def _attr_chart_b64(rows):
    """从逐场景 rows 生成 ΔPQ 归因图，返回 base64（失败返回 None）。"""
    if not rows:
        return None
    try:
        import io
        import base64
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import matplotlib.font_manager as fm
        for f in [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyhbd.ttc"]:
            if Path(f).exists():
                fm.fontManager.addfont(f)
        plt.rcParams["font.family"] = ["Microsoft YaHei", "DejaVu Sans"]
        plt.rcParams["axes.unicode_minus"] = False
        labels = [f"{r['noise_label']}\n{r['snr_db']}dB" for r in rows]
        deg  = [r["degraded_pq"] for r in rows]
        base = [r["baseline_pq"] for r in rows]
        cand = [r["candidate_pq"] for r in rows]
        fig, ax = plt.subplots(figsize=(8.6, 4.8), dpi=130)
        x = list(range(len(rows)))
        ax.plot(x, deg,  "--o", color="#94A3B8", lw=1.6, ms=4, label="退化语音")
        ax.plot(x, base, "-o", color="#7C3AED", lw=2.4, ms=5, label="已验收版本")
        ax.plot(x, cand, "-o", color="#0D9488", lw=2.4, ms=5, label="候选新版本")
        ax.set_xticks(x)
        ax.set_xticklabels(labels, fontsize=8.5)
        ax.set_ylabel("PQ 制作质量（越高越好）", fontsize=9)
        ax.set_title("版本迭代 · 逐噪声场景 PQ", fontsize=11, fontweight="bold")
        ax.grid(alpha=0.25, ls=":")
        ax.legend(fontsize=8.5, loc="upper right", framealpha=0.9)
        plt.tight_layout()
        buf = io.BytesIO()
        plt.savefig(buf, format="png", bbox_inches="tight")
        plt.close()
        return base64.b64encode(buf.getvalue()).decode()
    except Exception:
        return None


def llm_decide(context, config, rows=None):
    """真 LLM 决策（DeepSeek vision：看图 + 数字归因）。失败则回退规则。"""
    def _log(msg):
        try:
            with (LOOP_DIR / "llm_debug.log").open("a", encoding="utf-8") as f:
                f.write(datetime.datetime.now().strftime("%H:%M:%S") + " | " + msg + "\n")
        except Exception:
            pass
    _log(f"start key_len={len(DEEPSEEK_API_KEY)} model={DEEPSEEK_MODEL} rows={'Y' if rows else 'N'}")
    if not DEEPSEEK_API_KEY:
        _log("no api key -> return None")
        return None
    import requests

    constraints = context.get("allowed_actions") or {}
    prompt = (
        "你是通话降噪算法的「版本迭代」决策大脑（demo）。\n"
        "真实工作流：一个算法模型进来 → 迭代出新版本 → 看曲线是否整体上行 → 某场景还差就定向调参 → 出现「一个上、一个下」的权衡要诚实标注。\n"
        "只输出 JSON，不要任何多余文字。你是建议器，不能改写工程 gate。\n"
        f"本轮允许动作：{constraints.get('actions', [])}；iterate 可选版本：{constraints.get('iterate_versions', [])}。\n"
        f"已验收版本 baseline={config['baseline']}，候选新版本 candidate={config['candidate']}\n"
        f"可用版本候选：{[{'id': v['id'], 'strength': v['strength']} for v in list_versions()]}\n"
        "下面是逐场景 PQ 曲线图（紫=已验收，青=候选，灰虚线=退化语音）与数值：\n"
        + json.dumps(context, ensure_ascii=False, indent=2) + "\n"
        '输出格式：{"action": "...", "version": null, "hypothesis": "...", "reason": "..."}'
    )
    content = [{"type": "text", "text": prompt}]
    b64 = _attr_chart_b64(rows)
    _log(f"chart_b64_len={len(b64) if b64 else 'None'}")
    if b64:
        content.append({"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": b64}})
    try:
        resp = requests.post(
            f"{DEEPSEEK_BASE_URL.rstrip('/')}/v1/messages",
            headers={
                "x-api-key": DEEPSEEK_API_KEY,
                "anthropic-version": "2023-06-01",
                "content-type": "application/json",
            },
            json={
                "model": DEEPSEEK_MODEL,
                "max_tokens": 4000,
                "system": "你是严谨、诚实的通话降噪版本迭代决策者；看图 + 看数字做归因，不为了迭代而迭代，无解就诚实收口。",
                "messages": [{"role": "user", "content": content}],
            },
            timeout=120,
        )
        _log(f"http_status={resp.status_code}")
        resp.raise_for_status()
        data = resp.json()
        text = "".join(b.get("text", "") for b in data.get("content", []) if b.get("type") == "text")
        _log(f"text_len={len(text)} stop={data.get('stop_reason')}")
        m = re.search(r"\{.*\}", text, re.S)
        if not m:
            _log("no json in text -> return None")
            return None
        decision = json.loads(m.group(0))
        decision["llm"] = True
        decision["advisor"] = {
            "provider": "deepseek-anthropic-compatible",
            "model": DEEPSEEK_MODEL,
            "prompt_sha256": hashlib.sha256(prompt.encode("utf-8")).hexdigest(),
        }
        return decision
    except Exception as exc:
        try:
            (LOOP_DIR / "llm_error.log").write_text(
                f"{datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | {type(exc).__name__}: {exc}",
                encoding="utf-8",
            )
        except Exception:
            pass
        return None


def decide(judge_report, config, evaluated, rows=None):
    """LLM 只在工程 gate 允许的动作集合内给建议。"""
    allowed = allowed_actions(judge_report, evaluated)
    ctx = {
        "judge": judge_report,
        "baseline": config["baseline"],
        "candidate": config["candidate"],
        "version_pool": [{"id": v["id"], "strength": v["strength"]} for v in list_versions()],
        "evaluated": evaluated,
        "allowed_actions": allowed,
    }
    decision = llm_decide(ctx, config, rows)
    if decision is None:
        decision = rule_decide(judge_report, config, evaluated)
        decision["llm"] = False
    else:
        action = decision.get("action")
        invalid = action not in allowed["actions"]
        if action == "iterate" and decision.get("version") not in allowed["iterate_versions"]:
            invalid = True
        if invalid:
            decision = rule_decide(judge_report, config, evaluated)
            decision["llm"] = False
            decision["fallback_reason"] = "LLM 建议超出工程 gate 允许集合"
    decision["allowed_actions"] = allowed
    return decision


def build_blind_pairs(run_id, rows, include_assignments=False):
    """Create public blind pairs and keep identity-bearing assignments server-side."""
    pairs = []
    assignments = {}
    for row in rows or []:
        flip = int(hashlib.sha256(f"{run_id}:{row.get('stem')}".encode("utf-8")).hexdigest()[:2], 16) % 2 == 1
        base = {
            "role": "baseline", "model": row.get("baseline_model"), "file": row.get("baseline_file"),
            "path": row.get("baseline_path"),
        }
        cand = {
            "role": "candidate", "model": row.get("candidate_model"), "file": row.get("candidate_file"),
            "path": row.get("candidate_path"),
        }
        a, b = (cand, base) if flip else (base, cand)
        stem = row.get("stem")
        pairs.append({
            "stem": stem, "noise_label": row.get("noise_label"), "snr_db": row.get("snr_db"),
            "A": {"side": "A"}, "B": {"side": "B"},
        })
        assignments[stem] = {"A": a, "B": b}
    return (pairs, assignments) if include_assignments else pairs


def record_listening(run_id, pair, pick, note=""):
    if pick not in {"A", "B", "tie"}:
        raise ValueError("pick must be A, B, or tie")
    selected_role = "tie" if pick == "tie" else (pair.get(pick) or {}).get("role")
    stamp = datetime.datetime.now()
    token = hashlib.sha256(f"{run_id}:{pair.get('stem')}:{stamp.timestamp()}".encode("utf-8")).hexdigest()[:8]
    row = {
        "schema_version": 1,
        "listening_id": f"listen-{stamp.strftime('%Y%m%d-%H%M%S')}-{token}",
        "run_id": run_id,
        "timestamp": stamp.isoformat(timespec="seconds"),
        "stem": pair.get("stem"),
        "noise_label": pair.get("noise_label"),
        "snr_db": pair.get("snr_db"),
        "pick": pick,
        "selected_role": selected_role,
        "assignment": {"A": (pair.get("A") or {}).get("role"), "B": (pair.get("B") or {}).get("role")},
        "note": str(note or "")[:500],
    }
    LOOP_DIR.mkdir(parents=True, exist_ok=True)
    with _LISTENING_LOCK:
        with LISTENING_LOG.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    return row


def run_round(config, progress_cb=None):
    """跑一轮：补齐降噪 → 评测 → 归因 → 判定 → 决策建议。"""
    samples = collect_matrix(config, progress_cb=progress_cb)
    if not samples:
        return {"status": "error", "message": "矩阵里没有样本，请先在 04 生成对应退化 wav。"}
    baseline_v = get_version(config["baseline"])
    candidate_v = get_version(config["candidate"])
    if not baseline_v or not candidate_v:
        return {"status": "error", "message": "版本不存在，请检查 baseline / candidate。"}
    if baseline_v["id"] == candidate_v["id"]:
        return {"status": "error", "message": "候选新版本与已验收版本相同，请换一个版本。"}

    ensure_denoised(samples, baseline_v, progress_cb)
    ensure_denoised(samples, candidate_v, progress_cb)

    files = []
    for s in samples:
        files += [s["clean"], s["degraded"], str(_denoised_path(s, baseline_v)), str(_denoised_path(s, candidate_v))]
    unique = list(dict.fromkeys(files))
    if progress_cb:
        progress_cb(0, len(unique), "评测中…")
    if ((config.get("matrix") or {}).get("source") or "golden") == "golden":
        context = _golden_benchmark_context(config)
        scores, _ = _score_files_cached(unique, context["root"] / "scores.json", progress_cb=progress_cb)
    else:
        predictor = evaluate.load_predictor()
        scores = evaluate.score_files(predictor, unique, batch_size=4, progress_cb=progress_cb)

    rows = []
    base_pqs, cand_pqs, flags = [], [], []
    for s in samples:
        clean_path = str(Path(s["clean"]))
        degraded_path = str(Path(s["degraded"]))
        baseline_path = str(_denoised_path(s, baseline_v))
        candidate_path = str(_denoised_path(s, candidate_v))
        required = {
            "clean": scores.get(clean_path), "degraded": scores.get(degraded_path),
            "baseline": scores.get(baseline_path), "candidate": scores.get(candidate_path),
        }
        missing_scores = [name for name, value in required.items() if not value]
        if missing_scores:
            raise RuntimeError(f"评分矩阵不完整：{s.get('stem')} 缺少 {', '.join(missing_scores)}")
        deg, base, cand = required["degraded"], required["baseline"], required["candidate"]
        try:
            snr = float(s["snr_db"])
        except (TypeError, ValueError):
            raise RuntimeError(f"SNR 无效：{s.get('stem')} = {s.get('snr_db')!r}") from None
        if not math.isfinite(snr):
            raise RuntimeError(f"SNR 无效：{s.get('stem')} = {s.get('snr_db')!r}")
        is_regular = snr >= -5
        flags.append(is_regular)
        base_pqs.append(base["pq"])
        cand_pqs.append(cand["pq"])
        rows.append(
            {
                "scene_id": s["scene_id"],
                "noise_label": s["noise_label"],
                "snr_db": s["snr_db"],
                "degraded_pq": deg["pq"],
                "baseline_pq": base["pq"],
                "baseline_delta": round(base["pq"] - deg["pq"], 4),
                "candidate_pq": cand["pq"],
                "candidate_delta": round(cand["pq"] - deg["pq"], 4),
                "regular": is_regular,
                "stem": s["stem"],
                "clean_file": Path(s["clean"]).name,
                "degraded_file": Path(s["degraded"]).name,
                "baseline_file": _denoised_path(s, baseline_v).name,
                "candidate_file": _denoised_path(s, candidate_v).name,
                "baseline_path": baseline_path,
                "candidate_path": candidate_path,
                "baseline_model": baseline_v["model"],
                "candidate_model": candidate_v["model"],
                "voice_id": s.get("voice_id"),
                "voice_name": s.get("voice_name"),
                "cluster_id": s.get("cluster_id"),
                "cohort_id": s.get("cohort_id"),
            }
        )
    if len(rows) != len(samples):
        raise RuntimeError(f"完整矩阵应有 {len(samples)} 行，实际只得到 {len(rows)} 行评分。")

    judge_report = judge(base_pqs, cand_pqs, flags)
    # 按 SNR 降序（安静在前）
    rows.sort(key=lambda r: (-(float(r["snr_db"]) if r["snr_db"] not in ("", None) else -999.0)))
    return {
        "status": "ok",
        "baseline": baseline_v["id"],
        "candidate": candidate_v["id"],
        "baseline_model": baseline_v["model"],
        "candidate_model": candidate_v["model"],
        "baseline_strength": baseline_v["strength"],
        "candidate_strength": candidate_v["strength"],
        "rows": rows,
        "judge": judge_report,
    }


def _default_status():
    return {
        "schema_version": 3,
        "benchmark": json.loads(json.dumps(GOLDEN_BENCHMARK)),
        "config": json.loads(json.dumps(DEFAULT_CONFIG)),
        "evaluated_versions": [],
        "rounds": [],
        "changelog": [],
        "status": "idle",
        "workspace_mode": "blank",
        "baseline": None,
        "baseline_origin": None,
        "candidate": None,
        "last_judge": None,
        "last_decision": None,
        "conclusion": "",
    }


def load_status():
    with _STATUS_LOCK:
        if not LOOP_STATUS.exists():
            return _default_status()
        try:
            state = json.loads(LOOP_STATUS.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise RuntimeError(f"闭环状态文件损坏，已阻止用默认值覆盖：{LOOP_STATUS}") from exc
        if not isinstance(state, dict):
            raise RuntimeError(f"闭环状态文件格式无效，已阻止覆盖：{LOOP_STATUS}")
        state.setdefault("schema_version", 3)
        state.setdefault("benchmark", json.loads(json.dumps(GOLDEN_BENCHMARK)))
        state.setdefault("workspace_mode", "blank" if not state.get("baseline") else "project")
        state.setdefault("baseline_origin", _infer_baseline_origin(state))
        return state


def _infer_baseline_origin(state):
    baseline = state.get("baseline")
    if not baseline:
        return None
    accepted = next(
        (
            item for item in reversed(state.get("rounds") or [])
            if item.get("candidate") == baseline
            and (item.get("decision") or {}).get("action") in {"accept", "accept_conditional"}
        ),
        None,
    )
    if accepted:
        conditional = (accepted.get("decision") or {}).get("action") == "accept_conditional"
        return {
            "type": "conditional_promoted" if conditional else "promoted",
            "label": f"第 {accepted.get('round')} 轮确认条件晋级（极端回退已知）" if conditional else f"第 {accepted.get('round')} 轮确认晋级",
            "round": accepted.get("round"),
            "from_version": accepted.get("baseline"),
            "established_at": accepted.get("time"),
            "benchmark_id": (state.get("benchmark") or GOLDEN_BENCHMARK)["id"],
        }
    return {
        "type": "legacy",
        "label": "既有项目基准（来源待补充）",
        "established_at": state.get("updated_at"),
        "benchmark_id": (state.get("benchmark") or GOLDEN_BENCHMARK)["id"],
    }


def save_status(state):
    with _STATUS_LOCK:
        LOOP_STATUS.parent.mkdir(parents=True, exist_ok=True)
        state["updated_at"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        tmp = LOOP_STATUS.with_name(f".{LOOP_STATUS.name}.{uuid.uuid4().hex}.tmp")
        try:
            tmp.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(tmp, LOOP_STATUS)
        finally:
            if tmp.exists():
                tmp.unlink()
        return state


def _parse_history_time(value):
    """Parse the durable round timestamp without making archive policy fragile."""
    if isinstance(value, datetime.datetime):
        parsed = value
    else:
        text = str(value or "").strip()
        if not text:
            return None
        try:
            parsed = datetime.datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            parsed = None
            for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
                try:
                    parsed = datetime.datetime.strptime(text, fmt)
                    break
                except ValueError:
                    continue
            if parsed is None:
                return None
    return parsed.replace(tzinfo=None)


def _history_round_key(round_item):
    return ":".join(str(round_item.get(key) or "") for key in ("experiment_id", "round", "time", "baseline", "candidate"))


def _history_summary(round_item):
    """Keep the visible archive index compact; full rows remain in the local file."""
    return {
        key: round_item.get(key)
        for key in ("round", "experiment_id", "time", "baseline", "candidate", "decision", "judge")
        if key in round_item
    }


def _read_history_archive():
    if not HISTORY_ARCHIVE_FILE.is_file():
        return {"schema_version": 1, "rounds": []}
    try:
        payload = json.loads(HISTORY_ARCHIVE_FILE.read_text(encoding="utf-8"))
        return payload if isinstance(payload, dict) else {"schema_version": 1, "rounds": []}
    except (OSError, ValueError):
        return {"schema_version": 1, "rounds": []}


def _write_history_archive(payload):
    HISTORY_ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    tmp = HISTORY_ARCHIVE_FILE.with_name(f".{HISTORY_ARCHIVE_FILE.name}.{uuid.uuid4().hex}.tmp")
    try:
        tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(tmp, HISTORY_ARCHIVE_FILE)
    finally:
        if tmp.exists():
            tmp.unlink()


def _history_archive_label():
    try:
        return str(HISTORY_ARCHIVE_FILE.relative_to(PROJECT_ROOT))
    except ValueError:
        return str(HISTORY_ARCHIVE_FILE)


def prepare_history_archive(state, now=None):
    """Archive confirmed rounds older than a week while keeping the source ledger intact.

    The status ledger remains the source of truth for scoring and reports. The local
    archive is a durable retrieval copy, so hiding old rows in the UI never removes
    evidence needed to rebuild the V1–V4 summary.
    """
    now = now or datetime.datetime.now()
    cutoff = now - datetime.timedelta(days=HISTORY_ARCHIVE_AFTER_DAYS)
    current_period = now.strftime("%Y-%m")
    rounds = [item for item in (state.get("rounds") or []) if isinstance(item, dict)]
    archive = _read_history_archive()
    archived = [item for item in (archive.get("rounds") or []) if isinstance(item, dict)]
    known = {_history_round_key(item) for item in archived}
    stale = []
    newly_archived = 0
    for item in rounds:
        parsed = _parse_history_time(item.get("time"))
        if parsed and parsed < cutoff:
            stale.append(item)
            if _history_round_key(item) not in known:
                archived.append(item)
                known.add(_history_round_key(item))
                newly_archived += 1
    archived.sort(key=lambda item: _parse_history_time(item.get("time")) or datetime.datetime.min, reverse=True)
    if newly_archived:
        archive = {
            "schema_version": 1,
            "policy": {
                "archive_after_days": HISTORY_ARCHIVE_AFTER_DAYS,
                "storage": "local",
                "file": _history_archive_label(),
            },
            "updated_at": now.isoformat(timespec="seconds"),
            "rounds": archived,
        }
        _write_history_archive(archive)

    visible = []
    for item in rounds:
        parsed = _parse_history_time(item.get("time"))
        if parsed and parsed >= cutoff and parsed.strftime("%Y-%m") == current_period:
            visible.append(item)
    visible.sort(key=lambda item: _parse_history_time(item.get("time")) or datetime.datetime.min)
    return {
        "schema_version": 1,
        "current_period": current_period,
        "current_period_label": f"{now.year} 年 {now.month} 月",
        "archive_after_days": HISTORY_ARCHIVE_AFTER_DAYS,
        "storage": "local",
        "archive_file": _history_archive_label(),
        "visible_rounds": [item.get("round") for item in visible],
        "visible_count": len(visible),
        "archived_count": len(archived),
        "archived_rounds": [_history_summary(item) for item in archived],
        "newly_archived": newly_archived,
        "last_archived_at": archive.get("updated_at"),
    }


def load_history_archive():
    """Return the full local archive only when the user explicitly opens it."""
    archive = _read_history_archive()
    return {
        "schema_version": archive.get("schema_version", 1),
        "policy": archive.get("policy") or {
            "archive_after_days": HISTORY_ARCHIVE_AFTER_DAYS,
            "storage": "local",
            "file": _history_archive_label(),
        },
        "updated_at": archive.get("updated_at"),
        "rounds": archive.get("rounds") or [],
    }


@contextmanager
def status_transaction():
    """Serialize one read-modify-write transaction for loop_status.json."""
    with _STATUS_LOCK:
        state = load_status()
        yield state
        save_status(state)


def serialized_status_write(func):
    """Hold the status lock across an HTTP handler's read-modify-write cycle."""
    @wraps(func)
    def wrapped(*args, **kwargs):
        with _STATUS_LOCK:
            return func(*args, **kwargs)
    return wrapped


def start_new_experiment(state):
    """Clear only the active draft; keep reusable assets and confirmed history."""
    fresh = dict(state)
    fresh.update({
        "status": "idle",
        "baseline": None,
        "baseline_origin": None,
        "candidate": None,
        "last_judge": None,
        "last_decision": None,
        "conclusion": "",
        "workspace_mode": "blank",
        "workspace_started_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    })
    fresh["config"] = json.loads(json.dumps(DEFAULT_CONFIG))
    # evaluated_versions belongs to one experiment, while rounds/changelog are the
    # durable ledger. A blank experiment must not inherit the previous run's pool.
    fresh["evaluated_versions"] = []
    fresh["experiment_id"] = f"exp-{datetime.datetime.now().strftime('%Y%m%d-%H%M%S')}"
    return fresh


def _active_cluster_evidence():
    """Return the active C1-C4 assignment without importing the voice pipeline."""
    try:
        cohort_id = json.loads(ACTIVE_COHORT.read_text(encoding="utf-8")).get("cohort_id")
        result_path = VOICE_COHORTS_DIR / Path(str(cohort_id)).name / "cluster_result.json"
        result = json.loads(result_path.read_text(encoding="utf-8"))
        files = result.get("files") or []
        return {
            "cohort_id": cohort_id,
            "total_voices": len(files),
            "voice_to_cluster": {
                str(item.get("voice_name")): int(item.get("cluster")) + 1
                for item in files
                if item.get("voice_name") and item.get("cluster") is not None
            },
        }
    except Exception:
        return {"cohort_id": None, "total_voices": 0, "voice_to_cluster": {}}


def _voice_name_from_clean(filename):
    """Legacy exam names use `<voice name>_YYYYMMDD_HHMMSS.wav`."""
    stem = Path(str(filename or "")).stem
    return re.sub(r"_\d{8}_\d{6}$", "", stem)


def build_version_summary(state, preview=None):
    """Build the locked V1-V4 result model from shared measurements.

    Ranking uses fixed scene weights. Voices stay individually equal, so the
    anchored C1-C4 sizes become natural weights instead of four artificial 25%
    buckets. Disabled scenes are removed and the remaining weights renormalized.
    """
    versions = [item["id"] for item in VERSION_CANDIDATES]
    observations = {}
    runs = list(state.get("rounds") or [])
    if isinstance(preview, dict) and preview.get("status") == "ok":
        runs.append(preview)

    for run in runs:
        baseline = run.get("baseline")
        candidate = run.get("candidate")
        for row in run.get("rows") or []:
            canonical_scene = next(
                (scene for scene in channel.NOISE_SCENES if scene["id"] == row.get("scene_id")),
                None,
            )
            degraded_identity = row.get("degraded_file") or row.get("stem")
            if not degraded_identity:
                continue
            key = (
                row.get("cohort_id") or "legacy",
                row.get("voice_id") or row.get("clean_file") or "unknown-voice",
                row.get("scene_id") or row.get("noise_label") or "unknown-scene",
                str(row.get("snr_db")), degraded_identity,
            )
            item = observations.setdefault(key, {
                "scene_id": row.get("scene_id"),
                "noise_label": (canonical_scene or {}).get("label") or row.get("noise_label") or row.get("scene_id"),
                "snr_db": row.get("snr_db"),
                "regular": bool(row.get("regular", True)),
                "stem": row.get("stem"),
                "clean_file": row.get("clean_file"),
                "clean_pq": row.get("clean_pq"),
                "degraded_file": row.get("degraded_file"),
                "degraded_pq": row.get("degraded_pq"),
                "cohort_id": row.get("cohort_id"),
                "voice_id": row.get("voice_id"),
                "voice_name": row.get("voice_name"),
                "cluster_id": row.get("cluster_id"),
                "scores": {},
            })
            if baseline in versions and row.get("baseline_pq") is not None:
                item["scores"][baseline] = float(row["baseline_pq"])
            if candidate in versions and row.get("candidate_pq") is not None:
                item["scores"][candidate] = float(row["candidate_pq"])

    shared = [item for item in observations.values() if all(v in item["scores"] for v in versions)]

    def mean(values):
        return round(float(np.mean(values)), 4) if values else None

    metrics = []
    for version in versions:
        regular = [item["scores"][version] for item in shared if item["regular"]]
        extreme = [item["scores"][version] for item in shared if not item["regular"]]
        overall = regular + extreme
        metrics.append({
            "version": version,
            "coverage": len(overall),
            "regular_mean": mean(regular),
            "regular_worst": round(min(regular), 4) if regular else None,
            "extreme_mean": mean(extreme),
            "overall_mean": mean(overall),
        })

    ready = bool(observations) and len(shared) == len(observations)

    scene_groups = {}
    for item in shared:
        key = (item["noise_label"], item["snr_db"], item["regular"])
        group = scene_groups.setdefault(key, {version: [] for version in versions})
        for version in versions:
            group[version].append(item["scores"][version])
    scenes = [
        {
            "noise_label": key[0], "snr_db": key[1], "regular": key[2],
            "scores": {version: mean(values) for version, values in group.items()},
        }
        for key, group in scene_groups.items()
    ]
    scenes.sort(key=lambda item: -(float(item["snr_db"]) if item["snr_db"] not in (None, "") else -999))

    cluster_evidence = _active_cluster_evidence()
    cluster_groups = {cluster: {version: [] for version in versions} for cluster in range(1, 5)}
    mapped_voices = set()
    cluster_voices = {cluster: set() for cluster in range(1, 5)}
    for item in shared:
        voice_name = item.get("voice_name") or _voice_name_from_clean(item.get("clean_file"))
        cluster = item.get("cluster_id") or cluster_evidence["voice_to_cluster"].get(voice_name)
        if cluster not in cluster_groups:
            continue
        item["cluster_id"] = cluster
        mapped_voices.add(voice_name)
        cluster_voices[cluster].add(item.get("voice_id") or voice_name)
        for version in versions:
            cluster_groups[cluster][version].append(item["scores"][version])
    clusters = [
        {
            "cluster": f"C{cluster}",
            "sample_count": len(next(iter(group.values()))),
            "voice_count": len(cluster_voices[cluster]),
            "scores": {version: mean(values) for version, values in group.items()},
        }
        for cluster, group in cluster_groups.items()
    ]

    active_scene_ids = list(dict.fromkeys(
        item.get("scene_id") or item.get("noise_label") for item in shared
    ))
    configured = {scene_id: SCENE_WEIGHTS.get(scene_id) for scene_id in active_scene_ids}
    if active_scene_ids and all(configured[scene_id] is not None for scene_id in active_scene_ids):
        active_weight_total = sum(configured.values())
        effective = {
            scene_id: configured[scene_id] / active_weight_total
            for scene_id in active_scene_ids
        }
    elif active_scene_ids:
        # Legacy/custom matrices have no locked business weights; keep them usable
        # without pretending that their labels match the six-scene benchmark.
        effective = {scene_id: 1 / len(active_scene_ids) for scene_id in active_scene_ids}
    else:
        effective = {}

    scene_index = {}
    for item in shared:
        scene_id = item.get("scene_id") or item.get("noise_label")
        scene_index.setdefault(scene_id, []).append(item)

    weighted_scores = {}
    for version in versions:
        weighted_scores[version] = round(sum(
            effective[scene_id] * float(np.mean([item["scores"][version] for item in items]))
            for scene_id, items in scene_index.items()
        ), 4) if scene_index else None

    reference_version = state.get("baseline") if state.get("baseline") in versions else versions[0]
    reference_score = weighted_scores.get(reference_version)
    weighted_versions = []
    for version in versions:
        score = weighted_scores.get(version)
        weighted_versions.append({
            "version": version,
            "score": score,
            "delta": round(score - reference_score, 4)
            if score is not None and reference_score is not None else None,
        })

    scene_cluster_charts = []
    for scene_id in active_scene_ids:
        items = scene_index.get(scene_id) or []
        first = items[0] if items else {}
        chart_clusters = []
        for cluster in range(1, 5):
            cluster_items = [item for item in items if item.get("cluster_id") == cluster]
            reference_values = [item["scores"][reference_version] for item in cluster_items]
            reference_mean = float(np.mean(reference_values)) if reference_values else None
            values = []
            for version in versions:
                version_values = [item["scores"][version] for item in cluster_items]
                pq = float(np.mean(version_values)) if version_values else None
                values.append({
                    "version": version,
                    "pq": round(pq, 4) if pq is not None else None,
                    "delta": round(pq - reference_mean, 4)
                    if pq is not None and reference_mean is not None else None,
                })
            chart_clusters.append({
                "cluster": f"C{cluster}",
                "voice_count": len({
                    item.get("voice_id") or item.get("voice_name") or item.get("clean_file")
                    for item in cluster_items
                }),
                "values": values,
            })
        scene_cluster_charts.append({
            "scene_id": scene_id,
            "noise_label": first.get("noise_label") or scene_id,
            "snr_db": first.get("snr_db"),
            "configured_weight": configured.get(scene_id),
            "effective_weight": round(effective.get(scene_id, 0), 6),
            "clusters": chart_clusters,
        })

    recommendation = None
    if ready and weighted_versions:
        ranked = sorted(
            weighted_versions,
            key=lambda item: item["score"] if item["score"] is not None else -math.inf,
            reverse=True,
        )
        winner = ranked[0]
        risk_points = []
        for scene in scene_cluster_charts:
            for cluster in scene["clusters"]:
                value = next((v for v in cluster["values"] if v["version"] == winner["version"]), None)
                if value and value["delta"] is not None and value["delta"] < 0:
                    risk_points.append({
                        "scene_id": scene["scene_id"],
                        "noise_label": scene["noise_label"],
                        "cluster": cluster["cluster"],
                        "delta": value["delta"],
                    })
        risk_points.sort(key=lambda item: item["delta"])
        if winner["version"] == reference_version:
            risk_text = f"{reference_version.upper()} 仍是自然加权总分最高的基准版本。"
        elif risk_points:
            weakest = risk_points[0]
            risk_text = (
                f"主要风险在{weakest['noise_label']} · {weakest['cluster']}，"
                f"相对 {reference_version.upper()} 回退 {weakest['delta']:.3f} PQ。"
            )
        else:
            risk_text = f"六个场景的四个声线簇均未发现相对 {reference_version.upper()} 的均值回退。"
        recommendation = {
            "version": winner["version"],
            "policy": "六场景固定权重；每个音色等权，声线簇按 10/7/5/4 自然加权",
            "reason": (
                f"{winner['version'].upper()} 的自然加权 PQ 为 {winner['score']:.3f}，"
                f"相对 {reference_version.upper()} 为 {winner['delta']:+.3f}。"
            ),
            "risk": risk_text,
            "risk_points": risk_points,
        }

    total_unique_voices = len({
        item.get("voice_id") or item.get("voice_name") or item.get("clean_file")
        for item in shared
    })
    cluster_weights = [
        {
            "cluster": f"C{cluster}",
            "voice_count": len(cluster_voices[cluster]),
            "weight": round(len(cluster_voices[cluster]) / total_unique_voices, 6)
            if total_unique_voices else 0,
        }
        for cluster in range(1, 5)
    ]

    raw_evidence = sorted(shared, key=lambda item: (
        active_scene_ids.index(item.get("scene_id") or item.get("noise_label"))
        if (item.get("scene_id") or item.get("noise_label")) in active_scene_ids else 999,
        item.get("cluster_id") or 99,
        item.get("voice_name") or item.get("voice_id") or item.get("clean_file") or "",
    ))

    return {
        "status": "ready" if ready else "insufficient",
        "shared_samples": len(shared),
        "versions": metrics,
        "recommendation": recommendation,
        "scenes": scenes,
        "clusters": clusters,
        "reference_version": reference_version,
        "weighted_versions": weighted_versions,
        "scene_weights": [
            {
                "scene_id": scene_id,
                "configured_weight": configured.get(scene_id),
                "effective_weight": round(effective.get(scene_id, 0), 6),
            }
            for scene_id in active_scene_ids
        ],
        "cluster_weights": cluster_weights,
        "scene_cluster_charts": scene_cluster_charts,
        "raw_evidence": raw_evidence,
        "cluster_coverage": {
            "cohort_id": cluster_evidence["cohort_id"],
            "mapped_voices": len(mapped_voices),
            "total_voices": cluster_evidence["total_voices"],
            "complete": bool(cluster_evidence["total_voices"]) and len(mapped_voices) == cluster_evidence["total_voices"],
        },
    }


def _file_cache_signature(path):
    path = Path(path)
    stat = path.stat()
    return {"size": stat.st_size, "mtime_ns": stat.st_mtime_ns, "sha256": _sha256_file(path)}


def _evaluator_cache_signature():
    """Checkpoint content identity. Hashed only during scoring, not homepage polls."""
    payload = {
        "schema_version": ASSET_IDENTITY_SCHEMA,
        "impl_version": EVALUATE_IMPL_VERSION,
        "provider": "audiobox-aesthetics",
        "checkpoint": None,
        "sha256": None,
        "size": None,
        "mtime_ns": None,
    }
    checkpoint = evaluate.get_checkpoint()
    if not checkpoint:
        return payload
    path = Path(checkpoint)
    stat = path.stat()
    payload.update({
        "checkpoint": str(path.resolve()),
        "size": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
        "sha256": _sha256_file(path, memoize=False),
    })
    return payload


def _score_cache_entries(item):
    if not item:
        return []
    if isinstance(item.get("versions"), list):
        return list(item["versions"])
    if item.get("scores"):
        return [item]
    return []


def _lookup_cached_scores(item, signature, evaluator):
    for entry in _score_cache_entries(item):
        if entry.get("signature") == signature and entry.get("evaluator") == evaluator and entry.get("scores"):
            return entry["scores"]
    return None


def _store_cached_scores(cache, path, signature, evaluator, scores):
    versions = _score_cache_entries(cache.get(path))
    record = {"signature": signature, "evaluator": evaluator, "scores": scores}
    replaced = False
    for index, entry in enumerate(versions):
        if entry.get("evaluator") == evaluator and entry.get("signature") == signature:
            versions[index] = record
            replaced = True
            break
    if not replaced:
        versions.append(record)
    cache[path] = {"versions": versions}


def _score_files_cached(files, cache_path, progress_cb=None):
    """AudioBox cache keyed by audio content plus the evaluator asset identity."""
    cache = {}
    if cache_path.exists():
        try:
            cache = json.loads(cache_path.read_text(encoding="utf-8"))
        except Exception:
            cache = {}
    results = {}
    missing = []
    evaluator_signature = _evaluator_cache_signature()
    for raw in files:
        path = str(Path(raw))
        signature = _file_cache_signature(path)
        scores = _lookup_cached_scores(cache.get(path), signature, evaluator_signature)
        if scores:
            results[path] = scores
        else:
            missing.append(path)
    if missing:
        expected_signatures = {path: _file_cache_signature(path) for path in missing}
        if progress_cb:
            progress_cb(0, len(missing), f"加载 AudioBox · 待评测 {len(missing)} 条")
        predictor = evaluate.load_predictor()
        fresh = evaluate.score_files(predictor, missing, batch_size=4, progress_cb=progress_cb)
        absent = [path for path in missing if path not in fresh]
        if absent:
            raise RuntimeError(f"AudioBox 评分不完整，缺少 {len(absent)} 个结果：{absent[0]}")
        changed = [path for path in missing if _file_cache_signature(path) != expected_signatures[path]]
        if changed:
            raise RuntimeError(f"评分期间音频被改写，已拒绝写入缓存：{changed[0]}")
        for path in missing:
            scores = fresh[path]
            results[path] = scores
            _store_cached_scores(cache, path, expected_signatures[path], evaluator_signature, scores)
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        with _CACHE_LOCK:
            tmp = cache_path.with_name(f".{cache_path.name}.{uuid.uuid4().hex}.tmp")
            try:
                tmp.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
                os.replace(tmp, cache_path)
            finally:
                if tmp.exists():
                    tmp.unlink()
    elif progress_cb:
        progress_cb(len(files), len(files), "AudioBox 评分全部命中本地缓存")
    return results, len(missing)


def run_version_tournament(config=None, progress_cb=None):
    """Run all V1-V4 on the same 26 voices × fixed scenes benchmark."""
    effective = config or DEFAULT_CONFIG
    samples, context = prepare_golden_matrix(effective, progress_cb=progress_cb)
    versions = [get_version(item["id"]) for item in VERSION_CANDIDATES]
    for version in versions:
        ensure_denoised(samples, version, progress_cb=progress_cb)

    files = []
    for sample in samples:
        files.extend([sample["clean"], sample["degraded"]])
        files.extend(str(_denoised_path(sample, version)) for version in versions)
    unique = list(dict.fromkeys(files))
    scores, scored_now = _score_files_cached(unique, context["root"] / "scores.json", progress_cb=progress_cb)

    rows = []
    for sample in samples:
        clean = scores.get(str(Path(sample["clean"])))
        degraded = scores.get(str(Path(sample["degraded"])))
        version_scores = {
            version["id"]: (scores.get(str(_denoised_path(sample, version))) or {}).get("pq")
            for version in versions
        }
        if not clean or not degraded or any(value is None for value in version_scores.values()):
            continue
        snr = float(sample["snr_db"])
        rows.append({
            "scene_id": sample["scene_id"], "noise_label": sample["noise_label"],
            "snr_db": sample["snr_db"], "regular": snr >= -5,
            "stem": sample["stem"], "clean_file": Path(sample["clean"]).name,
            "clean_pq": clean["pq"],
            "degraded_file": Path(sample["degraded"]).name, "degraded_pq": degraded["pq"],
            "voice_id": sample["voice_id"], "voice_name": sample["voice_name"],
            "cluster_id": sample["cluster_id"], "cohort_id": sample["cohort_id"],
            "scores": version_scores,
        })
    if len(rows) != len(samples):
        raise RuntimeError(f"完整矩阵应有 {len(samples)} 行，实际只得到 {len(rows)} 行评分。")

    pseudo_rounds = []
    for version in versions[1:]:
        pseudo_rounds.append({
            "baseline": "v1", "candidate": version["id"],
            "rows": [{
                **row,
                "baseline_pq": row["scores"]["v1"],
                "candidate_pq": row["scores"][version["id"]],
            } for row in rows],
        })
    summary = build_version_summary({"baseline": "v1", "rounds": pseudo_rounds})
    result = {
        "status": "ok", "mode": "version_tournament", "schema_version": 1,
        "benchmark_id": GOLDEN_BENCHMARK["id"], "fingerprint": context["fingerprint"],
        "cohort_id": context["cohort_id"], "voices": len(context["voices"]),
        "scenes": len(context["scenes"]), "samples": len(rows),
        "scored_now": scored_now, "versions": [item["id"] for item in versions],
        "version_summary": summary, "rows": rows,
        "completed_at": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    result_path = context["root"] / "tournament_result.json"
    tmp = result_path.with_name(f".{result_path.name}.{uuid.uuid4().hex}.tmp")
    try:
        tmp.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(tmp, result_path)
    finally:
        if tmp.exists():
            tmp.unlink()
    result["result_path"] = str(result_path)
    return result


def load_version_tournament(config=None):
    try:
        context = _golden_benchmark_context(config or DEFAULT_CONFIG)
        path = context["root"] / "tournament_result.json"
        if not path.exists():
            return None
        result = json.loads(path.read_text(encoding="utf-8"))
        if result.get("fingerprint") != context["fingerprint"]:
            return None
        # Stored rows are the durable evidence. Rebuild derived presentation data
        # on read so UI/weighting fixes never require rerunning 806 audio scores.
        rows = result.get("rows") or []
        if rows:
            # Older saved tournaments did score the 26 clean voices but did not
            # copy clean PQ into every matrix row. Recover it from the score cache.
            try:
                score_cache = json.loads((context["root"] / "scores.json").read_text(encoding="utf-8"))
            except Exception:
                score_cache = {}
            for row in rows:
                if row.get("clean_pq") is not None:
                    continue
                clean_path = VOICE_COHORTS_DIR / Path(str(row.get("cohort_id") or "")).name / "audio" / Path(str(row.get("clean_file") or "")).name
                cached = score_cache.get(str(clean_path)) or {}
                entries = _score_cache_entries(cached)
                score = (entries[-1].get("scores") if entries else {}) or {}
                if score.get("pq") is not None:
                    row["clean_pq"] = score["pq"]
            pseudo_rounds = [
                {
                    "baseline": "v1",
                    "candidate": version["id"],
                    "rows": [
                        {
                            **row,
                            "baseline_pq": (row.get("scores") or {}).get("v1"),
                            "candidate_pq": (row.get("scores") or {}).get(version["id"]),
                        }
                        for row in rows
                    ],
                }
                for version in VERSION_CANDIDATES[1:]
            ]
            result["version_summary"] = build_version_summary({"baseline": "v1", "rounds": pseudo_rounds})
        result["result_path"] = str(path)
        return result
    except Exception:
        return None


def build_markdown(state):
    cfg = state.get("config") or {}
    lines = ["# 通话评测实验室 · 版本迭代报告", ""]
    lines.append(f"- 固定测试矩阵：{(state.get('benchmark') or GOLDEN_BENCHMARK).get('label')}")
    lines.append(f"- 当前基准版本（Baseline）：{state.get('baseline')}")
    lines.append(f"- 本轮待验证版本（Candidate）：{state.get('candidate')}")
    if state.get("baseline_origin"):
        lines.append(f"- 当前基准来源：{state['baseline_origin'].get('label')}")
    lines.append(f"- 状态：{state.get('status')}")
    if state.get("conclusion"):
        lines.append(f"- 结论：{state['conclusion']}")
    summary = state.get("version_summary") or build_version_summary(state)
    recommendation = summary.get("recommendation") or {}
    if recommendation:
        lines.extend([
            "",
            "## V1–V4 跨版本结论",
            f"- 工程推荐：**{recommendation.get('version')}**",
            f"- 选择规则：{recommendation.get('policy')}",
            f"- 原因：{recommendation.get('reason')}",
            f"- 公平比较样本：{summary.get('shared_samples')} 条（四版本均有评分）",
        ])
        coverage = summary.get("cluster_coverage") or {}
        if not coverage.get("complete"):
            lines.append(
                f"- C1–C4 证据缺口：当前仅映射 {coverage.get('mapped_voices', 0)}/{coverage.get('total_voices', 0)} 个固定音色，簇级结果不参与推荐。"
            )
    lines.append("")
    for r in state.get("rounds", []):
        j = r.get("judge") or {}
        d = r.get("decision") or {}
        lines.append(f"## 第 {r.get('round')} 轮 · {r.get('baseline')} vs {r.get('candidate')}")
        lines.append(f"- 实验：{r.get('experiment_id') or '未记录'} · Run：{r.get('run_id') or '未记录'} · 时间：{r.get('time') or '未记录'}")
        lines.append(f"- 判定：**{j.get('verdict')}**（常规 Δ={j.get('regular_delta')}，p={j.get('regular_p')}，极端 Δ={j.get('extreme_delta')}）")
        if d:
            lines.append(f"- 建议：{d.get('action')} {d.get('model') or ''} —— {d.get('hypothesis')}")
        rows = r.get("rows") or []
        if rows:
            def _score(value):
                try:
                    return float(value) if value is not None and not isinstance(value, bool) and math.isfinite(float(value)) else None
                except (ValueError, TypeError):
                    return None

            def _num(value):
                value = _score(value)
                if value is None:
                    return "—"
                return f"{float(value):.4f}".rstrip("0").rstrip(".")

            def _cell(value):
                return str(value).replace("|", "\\|").replace("\r", " ").replace("\n", " ")

            lines.extend([
                "",
                "### 本轮逐样本证据",
                "| 场景 | 音色 | 样本 ID | SNR | 当前版本 PQ | 待验证版本 PQ | 新版−当前 ΔPQ | 判读 |",
                "| --- | --- | --- | ---: | ---: | ---: | ---: | --- |",
            ])
            for row in rows:
                baseline_pq = _score(row.get("baseline_pq"))
                candidate_pq = _score(row.get("candidate_pq"))
                delta = round(float(candidate_pq) - float(baseline_pq), 4) if baseline_pq is not None and candidate_pq is not None else None
                delta_text = "—" if delta is None else f"{delta:+.4f}".rstrip("0").rstrip(".")
                verdict = "无可比较数据" if delta is None else "上行" if delta >= 0 else "回退"
                if row.get("regular") is False:
                    verdict += " · 极端"
                lines.append(
                    f"| {_cell(row.get('noise_label') or row.get('scene_id') or '—')} | {_cell(row.get('voice_name') or row.get('voice_id') or '音色未记录')} | {_cell(row.get('stem') or '样本 ID 未记录')} | {_cell(row.get('snr_db', '—'))} dB | {_num(baseline_pq)} | {_num(candidate_pq)} | {delta_text} | {verdict} |"
                )
        lines.append("")
    return "\n".join(lines)


def export_report_excel(state, out_path):
    """Export the decision ledger and full sample evidence as an Excel workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "概览"
    benchmark = state.get("benchmark") or GOLDEN_BENCHMARK
    summary = state.get("version_summary") or build_version_summary(state)
    recommendation = summary.get("recommendation") or {}
    overview_rows = [
        ("固定测试矩阵", benchmark.get("label")),
        ("当前基准版本", state.get("baseline")),
        ("待验证版本", state.get("candidate")),
        ("状态", state.get("status")),
        ("结论", state.get("conclusion")),
        ("跨版本工程推荐", recommendation.get("version")),
        ("推荐原因", recommendation.get("reason")),
        ("公平比较样本", summary.get("shared_samples")),
        ("已确认轮次", len(state.get("rounds") or [])),
    ]
    overview.append(["字段", "内容"])
    for row in overview_rows:
        overview.append(list(row))

    rounds_sheet = workbook.create_sheet("轮次")
    rounds_sheet.append(["轮次", "实验 ID", "Run ID", "时间", "基准版本", "候选版本", "常规平均 ΔPQ", "常规 p 值", "极端平均 ΔPQ", "工程判定", "处理动作", "说明"])
    samples_sheet = workbook.create_sheet("样本明细")
    samples_sheet.append(["轮次", "实验 ID", "Run ID", "场景", "参考环境声级（约 dB，非实测）", "仿真 SNR dB", "音色", "样本 ID", "当前版本 PQ", "待验证版本 PQ", "新版−当前 ΔPQ", "判读"])
    scene_by_id = {item["id"]: item for item in channel.NOISE_SCENES}
    scene_by_label = {item["label"]: item for item in channel.NOISE_SCENES}

    def number(value):
        try:
            parsed = float(value)
            return parsed if math.isfinite(parsed) else None
        except (TypeError, ValueError):
            return None

    for round_item in state.get("rounds") or []:
        judge = round_item.get("judge") or {}
        decision = round_item.get("decision") or {}
        rounds_sheet.append([
            round_item.get("round"), round_item.get("experiment_id"), round_item.get("run_id"), round_item.get("time"),
            round_item.get("baseline"), round_item.get("candidate"), number(judge.get("regular_delta")),
            number(judge.get("regular_p")), number(judge.get("extreme_delta")), judge.get("verdict"),
            decision.get("action"), decision.get("reason") or decision.get("hypothesis"),
        ])
        for row in round_item.get("rows") or []:
            baseline_pq = number(row.get("baseline_pq"))
            candidate_pq = number(row.get("candidate_pq"))
            delta = candidate_pq - baseline_pq if baseline_pq is not None and candidate_pq is not None else None
            scene = scene_by_id.get(row.get("scene_id")) or scene_by_label.get(row.get("noise_label")) or {}
            verdict = "无可比较数据" if delta is None else "上行" if delta >= 0 else "回退"
            if row.get("regular") is False:
                verdict += " · 极端"
            samples_sheet.append([
                round_item.get("round"), round_item.get("experiment_id"), round_item.get("run_id"),
                row.get("noise_label") or row.get("scene_id"), scene.get("level_db"), row.get("snr_db"),
                row.get("voice_name") or row.get("voice_id"), row.get("stem"), baseline_pq, candidate_pq, delta, verdict,
            ])

    header_fill = PatternFill("solid", fgColor="252A27")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 48)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 12)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(out_path)
    return out_path
